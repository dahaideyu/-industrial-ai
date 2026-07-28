# cython: annotation_typing=False, infer_types=False, language_level=3
"""文件转换任务：Word/PPT转PDF、Excel转HTML、Markdown转HTML"""
import io
import logging
import os
import tempfile
import re

from backend.services.knowledge_management.tasks.celery_app import kb_celery
from backend.core.knowledge_management.database import SessionLocal
from backend.core.knowledge_management.models import DocumentVersion
from backend.clients.knowledge_management.libreoffice_client import libreoffice_client
from backend.clients.knowledge_management.minio_client import minio_client

logger = logging.getLogger(__name__)

# Excel 转 HTML 的扩展名
_EXCEL_EXTS = {".xls", ".xlsx", ".xlsm"}


def _excel_to_html(file_path: str, file_name: str) -> str:
    """将 Excel 文件转为 HTML 字符串。"""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    wb = openpyxl.load_workbook(file_path, data_only=True)
    html_parts = ['<html><head><meta charset="utf-8"><style>',
                  '*{box-sizing:border-box;}',
                  'body{font-family:Arial,sans-serif;margin:0;padding:0;overflow:auto;display:flex;flex-direction:column;height:100vh;}',
                  'table{border-collapse:collapse;font-size:12px;width:max-content;}',
                  'td,th{border:1px solid #c0c0c0;padding:2px 6px;white-space:nowrap;}',
                  'th{background:#e8e8e8;font-weight:600;}',
                  '.sheet-nav{display:flex;gap:4px;padding:6px 12px;background:#f0f0f0;border-bottom:2px solid #d0d0d0;flex-shrink:0;}',
                  '.sheet-nav button{border:1px solid #c0c0c0;background:#fff;padding:4px 12px;cursor:pointer;font-size:12px;border-radius:3px;}',
                  '.sheet-nav button.active{background:#2b7c3b;color:#fff;border-color:#2b7c3b;}',
                  '.sheet-content{flex:1;overflow:auto;padding:8px;}',
                  '.sheet-page{display:none;}',
                  '.sheet-page.active{display:block;}',
                  '</style></head><body>']

    sheet_names = wb.sheetnames

    if len(sheet_names) > 1:
        html_parts.append('<div class="sheet-nav">')
        for i, name in enumerate(sheet_names):
            active = ' active' if i == 0 else ''
            html_parts.append(f'<button class="sheet-btn{active}" onclick="switchSheet({i})">{name}</button>')
        html_parts.append('</div>')

    html_parts.append('<div class="sheet-content">')

    for si, sheet_name in enumerate(sheet_names):
        ws = wb[sheet_name]
        show = ' active' if si == 0 else ''
        html_parts.append(f'<div class="sheet-page{show}" id="sheet-{si}">')
        html_parts.append('<table>')

        # 确定有效范围
        if ws.max_row and ws.max_column:
            # 表头
            html_parts.append('<tr>')
            for col in range(1, ws.max_column + 1):
                html_parts.append(f'<th>{get_column_letter(col)}</th>')
            html_parts.append('</tr>')

            # 数据行
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
                html_parts.append('<tr>')
                for cell in row:
                    val = cell.value
                    if val is None:
                        val = ''
                    elif isinstance(val, float):
                        val = str(val) if val == int(val) else f'{val:g}'
                    else:
                        val = str(val)

                    # 应用样式
                    style = ''
                    if cell.fill and cell.fill.fgColor:
                        try:
                            rgb = str(cell.fill.fgColor.rgb)
                            if rgb and rgb != '00000000':
                                style += f'background-color:#{rgb[2:]};'
                        except Exception:
                            pass
                    if cell.font:
                        if cell.font.bold:
                            style += 'font-weight:bold;'
                        if cell.font.color:
                            try:
                                rgb = str(cell.font.color.rgb)
                                if rgb and rgb != '00000000':
                                    style += f'color:#{rgb[2:]};'
                            except Exception:
                                pass
                    if cell.alignment and cell.alignment.horizontal:
                        style += f'text-align:{cell.alignment.horizontal};'
                    if cell.alignment and cell.alignment.vertical:
                        style += f'vertical-align:{cell.alignment.vertical};'

                    # 合并单元格处理
                    colspan = ''
                    rowspan = ''
                    for merged in ws.merged_cells.ranges:
                        if cell.coordinate in merged:
                            if merged.min_col != merged.max_col:
                                colspan = f' colspan="{merged.max_col - merged.min_col + 1}"'
                            if merged.min_row != merged.max_row:
                                rowspan = f' rowspan="{merged.max_row - merged.min_row + 1}"'
                            break

                    style_attr = f' style="{style}"' if style else ''
                    html_parts.append(f'<td{style_attr}{colspan}{rowspan}>{val}</td>')
                html_parts.append('</tr>')

        html_parts.append('</table>')
        html_parts.append('</div>')

    html_parts.append('</div>')

    if len(sheet_names) > 1:
        html_parts.append('<script>function switchSheet(i){'
                         'document.querySelectorAll(".sheet-page").forEach(p=>p.classList.remove("active"));'
                         'document.getElementById("sheet-"+i).classList.add("active");'
                         'document.querySelectorAll(".sheet-btn").forEach(b=>b.classList.remove("active"));'
                         'document.querySelector(".sheet-btn:nth-child("+(i+1)+")").classList.add("active");'
                         '}</script>')

    html_parts.append('</body></html>')
    return '\n'.join(html_parts)


@kb_celery.task(bind=True, queue="kb_convert")
def convert_to_pdf(self, version_id: str):
    """将 Word/PPT 文档转换为 PDF 预览。"""
    db = SessionLocal()
    try:
        version = db.query(DocumentVersion).filter(DocumentVersion.id == version_id).first()
        if not version:
            return

        version.convert_status = "processing"
        db.commit()

        tmp_dir = tempfile.mkdtemp(prefix="knb_convert_")
        local_path = os.path.join(tmp_dir, version.original_filename)
        minio_client.download_file(version.storage_path, local_path)

        pdf_path = libreoffice_client.convert_to_pdf(local_path, tmp_dir)
        pdf_object = f"{version.storage_path.rsplit('.', 1)[0]}_preview.pdf"
        minio_client.upload_file(pdf_object, pdf_path, "application/pdf")

        version.pdf_preview_path = pdf_object
        version.convert_status = "done"
        db.commit()
        logger.info("转 PDF 完成: version_id=%s", version_id)
    except Exception as exc:
        logger.exception("转 PDF 失败: version_id=%s", version_id)
        try:
            version.convert_status = "failed"
            db.commit()
        except Exception:
            db.rollback()
        raise self.retry(exc=exc)
    finally:
        db.close()


@kb_celery.task(bind=True, queue="kb_convert")
def convert_excel_to_html(self, version_id: str):
    """将 Excel 文件转换为 HTML 预览。"""
    db = SessionLocal()
    try:
        version = db.query(DocumentVersion).filter(DocumentVersion.id == version_id).first()
        if not version:
            return

        version.convert_status = "processing"
        db.commit()

        tmp_dir = tempfile.mkdtemp(prefix="knb_excel_")
        local_path = os.path.join(tmp_dir, version.original_filename)
        minio_client.download_file(version.storage_path, local_path)

        html_content = _excel_to_html(local_path, version.original_filename)
        html_object = f"{version.storage_path.rsplit('.', 1)[0]}_preview.html"
        minio_client.upload_fileobj(html_object, io.BytesIO(html_content.encode('utf-8')), len(html_content.encode('utf-8')), "text/html")

        version.pdf_preview_path = html_object
        version.convert_status = "done"
        db.commit()
        logger.info("Excel 转 HTML 完成: version_id=%s", version_id)
    except Exception as exc:
        logger.exception("Excel 转 HTML 失败: version_id=%s", version_id)
        try:
            version.convert_status = "failed"
            db.commit()
        except Exception:
            db.rollback()
        raise
    finally:
        db.close()


def _md_to_html_body(md_text: str) -> str:
    """将 Markdown 文本转换为 HTML body 内容。

    处理流程：保护代码块 → 转换 Markdown 语法 → 还原代码块 → 包裹 HTML 模板。
    """
    text = md_text

    # ---- 保护代码块和行内代码（避免被后续正则误处理） ----
    code_blocks = {}
    inline_codes = {}

    # 代码块 (```...```)
    def _protect_code_blocks(m):
        nonlocal code_blocks
        lang = m.group(1) or ""
        code = m.group(2)
        placeholder = f"\x00CB{len(code_blocks)}\x00"
        lang_attr = f' class="language-{lang}"' if lang else ""
        code_blocks[placeholder] = f"<pre><code{lang_attr}>{code}</code></pre>"
        return placeholder
    text = re.sub(r"```(\w*)\n(.*?)```", _protect_code_blocks, text, flags=re.DOTALL)

    # 行内代码 (`...`)
    def _protect_inline_codes(m):
        nonlocal inline_codes
        placeholder = f"\x00IC{len(inline_codes)}\x00"
        inline_codes[placeholder] = f"<code>{m.group(1)}</code>"
        return placeholder
    text = re.sub(r"`([^`]+)`", _protect_inline_codes, text)

    # ---- Markdown → HTML 转换 ----
    # 标题（h4→h1，先匹配更具体的）
    text = re.sub(r"^#### (.+)$", r"<h4>\1</h4>", text, flags=re.MULTILINE)
    text = re.sub(r"^### (.+)$", r"<h3>\1</h3>", text, flags=re.MULTILINE)
    text = re.sub(r"^## (.+)$", r"<h2>\1</h2>", text, flags=re.MULTILINE)
    text = re.sub(r"^# (.+)$", r"<h1>\1</h1>", text, flags=re.MULTILINE)

    # 加粗和斜体
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"\*(.+?)\*", r"<em>\1</em>", text)

    # 图片（必须在链接之前，避免 ![...](...) 被误匹配为链接）
    text = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", r'<img src="\2" alt="\1">', text)

    # 链接
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)

    # 分割线
    text = re.sub(r"^---+$", r"<hr>", text, flags=re.MULTILINE)

    # 无序列表
    text = re.sub(r"^- (.+)$", r"<li>\1</li>", text, flags=re.MULTILINE)

    # 引用块
    text = re.sub(r"^> (.+)$", r"<blockquote>\1</blockquote>", text, flags=re.MULTILINE)

    # 段落：连续两个换行 → 段落分隔
    text = re.sub(r"\n\n+", "\n<p>\n", text)
    # 单个换行 → <br>（但不影响已生成的 HTML 标签）
    text = re.sub(r"\n(?!<)", "<br>\n", text)

    # 还原代码块和行内代码
    for placeholder, html in code_blocks.items():
        text = text.replace(placeholder, html)
    for placeholder, html in inline_codes.items():
        text = text.replace(placeholder, html)

    # 包裹完整 HTML
    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<style>
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif;
    line-height: 1.7; color: #1f2937; max-width: 900px; margin: 0 auto; padding: 24px 32px;
  }}
  h1 {{ font-size: 1.8em; border-bottom: 2px solid #e5e7eb; padding-bottom: 8px; margin: 24px 0 12px; }}
  h2 {{ font-size: 1.5em; border-bottom: 1px solid #f3f4f6; padding-bottom: 6px; margin: 20px 0 10px; }}
  h3 {{ font-size: 1.25em; margin: 16px 0 8px; }}
  h4 {{ font-size: 1.1em; margin: 12px 0 6px; }}
  p {{ margin: 8px 0; }}
  pre {{ background: #f3f4f6; border: 1px solid #e5e7eb; border-radius: 6px; padding: 12px 16px; overflow-x: auto; font-size: 13px; line-height: 1.5; }}
  code {{ background: #f3f4f6; padding: 1px 5px; border-radius: 3px; font-size: 0.9em; font-family: "Fira Code", "Consolas", monospace; }}
  pre code {{ background: none; padding: 0; }}
  li {{ margin: 4px 0; }}
  a {{ color: #2563eb; text-decoration: none; }}
  a:hover {{ text-decoration: underline; }}
  hr {{ border: none; border-top: 1px solid #e5e7eb; margin: 20px 0; }}
  img {{ max-width: 100%; border-radius: 4px; }}
  blockquote {{ border-left: 3px solid #d1d5db; padding-left: 16px; color: #6b7280; margin: 12px 0; }}
  table {{ border-collapse: collapse; width: 100%; margin: 12px 0; }}
  th, td {{ border: 1px solid #d1d5db; padding: 6px 12px; text-align: left; }}
  th {{ background: #f9fafb; font-weight: 600; }}
</style>
</head>
<body>
{text}
</body>
</html>"""
    return html


@kb_celery.task(bind=True, queue="kb_convert")
def convert_md_to_html(self, version_id: str):
    """将 Markdown 文件转换为 HTML 预览页面。"""
    db = SessionLocal()
    try:
        version = db.query(DocumentVersion).filter(DocumentVersion.id == version_id).first()
        if not version:
            return

        version.convert_status = "processing"
        db.commit()

        tmp_dir = tempfile.mkdtemp(prefix="knb_md_")
        local_path = os.path.join(tmp_dir, version.original_filename)
        minio_client.download_file(version.storage_path, local_path)

        with open(local_path, "r", encoding="utf-8") as f:
            md_text = f.read()

        html_content = _md_to_html_body(md_text)
        html_object = f"{version.storage_path.rsplit('.', 1)[0]}_preview.html"
        html_bytes = html_content.encode("utf-8")
        minio_client.upload_fileobj(
            html_object,
            io.BytesIO(html_bytes),
            len(html_bytes),
            "text/html",
        )

        version.pdf_preview_path = html_object
        version.convert_status = "done"
        db.commit()
        logger.info("Markdown 转 HTML 完成: version_id=%s", version_id)
    except Exception as exc:
        logger.exception("Markdown 转 HTML 失败: version_id=%s", version_id)
        try:
            version.convert_status = "failed"
            db.commit()
        except Exception:
            db.rollback()
        raise
    finally:
        db.close()
